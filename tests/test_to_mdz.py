"""convert_to_mdz and the markitdown-mdz command."""

from __future__ import annotations

import base64
import io
import os
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout

import mdzip_core
from markitdown import StreamInfo

from markitdown_mdzip import convert_to_mdz
from markitdown_mdzip.__main__ import main

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)
PNG2 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFBQIAX8jx0gAAAABJRU5ErkJggg=="
)


def uri(data: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(data).decode()


def html_with_images() -> bytes:
    return (
        "<html><head><title>Trip Report</title></head><body><h1>Trip Report</h1>"
        f'<p><img src="{uri(PNG)}" alt="map"></p>'
        f'<p><img src="{uri(PNG2)}" alt="photo"></p>'
        f'<p><img src="{uri(PNG)}" alt="map again"></p>'
        "</body></html>"
    ).encode()


def optional(module: str):
    try:
        __import__(module)
        return None
    except ImportError:
        return f"{module} not installed"


class HtmlTests(unittest.TestCase):
    def test_images_become_archive_files(self):
        result = convert_to_mdz(io.BytesIO(html_with_images()), stream_info=StreamInfo(extension=".html", filename="trip.html"))
        self.assertEqual(result.entry_point, "trip.md")
        self.assertEqual(result.title, "Trip Report")
        self.assertEqual(result.images, ["images/image-001.png", "images/image-002.png"])
        self.assertNotIn("data:", result.markdown)
        self.assertIn("![map](images/image-001.png)", result.markdown)
        self.assertIn("![photo](images/image-002.png)", result.markdown)
        self.assertIn("![map again](images/image-001.png)", result.markdown)  # stored once

        with mdzip_core.open(result.data) as archive:
            self.assertEqual(archive.entry_point, "trip.md")
            self.assertEqual(archive.title, "Trip Report")
            self.assertEqual(archive.read_bytes("images/image-001.png"), PNG)
            self.assertEqual(archive.read_bytes("images/image-002.png"), PNG2)
            self.assertEqual(archive.resolve_link("trip.md", "images/image-002.png").kind, "internal")
            self.assertEqual([d for d in archive.validate() if d.level == "error"], [])

    def test_from_a_path_and_options(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "Trip Report.html")
            with open(path, "wb") as f:
                f.write(html_with_images())
            result = convert_to_mdz(path, title="Custom", images_dir="media/")
        self.assertEqual(result.entry_point, "Trip Report.md")
        self.assertEqual(result.title, "Custom")
        self.assertEqual(result.images[0], "media/image-001.png")

    def test_entry_names(self):
        page = b"<html><body><p>Hi</p></body></html>"
        for filename, expected in [("README.html", "index.md"), ("", "index.md"), ("notes.html", "notes.md")]:
            result = convert_to_mdz(io.BytesIO(page), stream_info=StreamInfo(extension=".html", filename=filename))
            self.assertEqual(result.entry_point, expected, filename)
        explicit = convert_to_mdz(io.BytesIO(page), stream_info=StreamInfo(extension=".html"), entry_point="doc.md")
        self.assertEqual(explicit.entry_point, "doc.md")


@unittest.skipIf(optional("docx") or optional("mammoth"), "python-docx and markitdown[docx] needed")
class DocxTests(unittest.TestCase):
    def test_word_images_are_kept(self):
        import docx

        document = docx.Document()
        document.add_heading("Quarterly Report", level=1)
        document.add_paragraph("Sales went up.")
        document.add_picture(io.BytesIO(PNG))
        buf = io.BytesIO()
        document.save(buf)
        buf.seek(0)

        result = convert_to_mdz(buf, stream_info=StreamInfo(extension=".docx", filename="report.docx"))
        self.assertEqual(result.entry_point, "report.md")
        self.assertIn("# Quarterly Report", result.markdown)
        self.assertEqual(result.title, "Quarterly Report")  # Word has no title; the heading is used
        self.assertEqual(result.images, ["images/image-001.png"])
        with mdzip_core.open(result.data) as archive:
            self.assertEqual(archive.read_bytes("images/image-001.png"), PNG)


@unittest.skipIf(optional("pptx"), "markitdown[pptx] needed")
class PptxTests(unittest.TestCase):
    def test_slide_images_are_kept(self):
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Roadmap"
        slide.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(1))
        buf = io.BytesIO()
        deck.save(buf)
        buf.seek(0)

        result = convert_to_mdz(buf, stream_info=StreamInfo(extension=".pptx", filename="deck.pptx"))
        self.assertIn("Roadmap", result.markdown)
        self.assertEqual(result.images, ["images/image-001.png"])
        self.assertIn("](images/image-001.png)", result.markdown)


class CommandTests(unittest.TestCase):
    def run_main(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            try:
                code = main(list(args))
            except SystemExit as exc:
                code = exc.code
        return code, out.getvalue(), err.getvalue()

    def test_writes_next_to_the_source_and_refuses_to_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = os.path.join(tmp, "trip.html")
            with open(source, "wb") as f:
                f.write(html_with_images())
            output = os.path.join(tmp, "trip.mdz")

            code, out, _ = self.run_main(source, "-o", output)
            self.assertEqual(code, 0)
            self.assertIn("trip.md + 2 images", out)
            with mdzip_core.open(output) as archive:
                self.assertEqual(archive.entry_point, "trip.md")

            code, _, err = self.run_main(source, "-o", output)
            self.assertNotEqual(code, 0)
            self.assertIn("already exists", err)
            self.assertEqual(self.run_main(source, "-o", output, "--force")[0], 0)

    def test_missing_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            code, _, err = self.run_main(os.path.join(tmp, "nope.docx"), "-o", os.path.join(tmp, "out.mdz"))
        self.assertEqual(code, 1)
        self.assertIn("markitdown-mdz:", err)


if __name__ == "__main__":
    unittest.main()
